import pypdf
import docx
import openpyxl

def extract_text(file_obj, filename):
    name = filename.lower()
    if name.endswith('.pdf'):
        return _extract_pdf(file_obj)
    elif name.endswith('.docx'):
        return _extract_docx(file_obj)
    elif name.endswith('.xlsx') or name.endswith('.xls'):
        return _extract_excel(file_obj)
    elif name.endswith('.txt'):
        return file_obj.read().decode('utf-8', errors='ignore')
    else:
        raise ValueError(f"Unsupported file type: {filename}")
    
def _extract_pdf(file_obj):
    reader = pypdf.PdfReader(file_obj)
    parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            parts.append(text)
    return '\n'.join(parts)

def _extract_docx(file_obj):
    doc = docx.Document(file_obj)
    return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())

def _extract_excel(file_obj):
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"Sheet - {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            row_text = '\t'.join(str(c) if c is not None else '' for c in row)
            if row_text.strip():
                parts.append(row_text)
    return '\n'.join(parts)

def chunk_text(text, chunk_size=300, overlap=50):
    words = text.split()
    chunks = []
    start = 0
    while start < len(words):
        end = start + chunk_size
        chunk = ' '.join(words[start:end])
        if chunk.strip():
            chunks.append(chunk)
        start += chunk_size - overlap
    return chunks